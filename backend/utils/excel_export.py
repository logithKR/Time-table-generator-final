import io
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
from sqlalchemy.orm import Session
import models

def generate_timetable_excel_bytes(department_code: str, semester: int, db: Session) -> io.BytesIO:
    slots = db.query(models.SlotMaster).order_by(models.SlotMaster.day_of_week, models.SlotMaster.period_number).all()
    if not slots:
        raise ValueError("No slots configured")

    days_order = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
    active_days = [d for d in days_order if any(s.day_of_week == d for s in slots)]
    if not active_days: active_days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    max_period = max([s.period_number for s in slots], default=8)
    periods = list(range(1, max_period + 1))

    query = db.query(models.TimetableEntry)
    if department_code: query = query.filter(models.TimetableEntry.department_code == department_code)
    if semester: query = query.filter(models.TimetableEntry.semester == semester)
    entries = query.all()

    timetable_dict = {}
    for e in entries:
        key = (e.day_of_week, e.period_number)
        timetable_dict.setdefault(key, []).append(e)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Timetable {department_code or 'All'} Sem {semester or 'All'}"

    thin = Side(border_style="thin", color="D1D5DB")
    border_all = Border(top=thin, left=thin, right=thin, bottom=thin)
    fill_header = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_day = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_lab = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_class = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_empty = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    font_bold = Font(bold=True, color="1F2937", size=11)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    header_row = 2
    d_cell = ws.cell(row=header_row, column=1, value="DAY / PERIOD")
    d_cell.font = font_bold
    d_cell.alignment = align_center
    d_cell.fill = fill_day
    d_cell.border = border_all
    ws.column_dimensions['A'].width = 16

    for i, p in enumerate(periods):
        col = i + 2
        slot = next((s for s in slots if s.period_number == p and s.day_of_week == 'Monday'), None)
        time_str = f"{slot.start_time} - {slot.end_time}" if slot else ""
        header_text = f"Period {p}\n({time_str})"
        c = ws.cell(row=header_row, column=col, value=header_text)
        c.font = Font(bold=True, color="374151", size=10)
        c.alignment = align_center
        c.fill = fill_header
        c.border = border_all
        ws.column_dimensions[get_column_letter(col)].width = 24

    current_row = header_row + 1
    for day in active_days:
        dc = ws.cell(row=current_row, column=1, value=day[:3].upper())
        dc.font = Font(bold=True, size=12, color="1F2937")
        dc.alignment = align_center
        dc.fill = fill_day
        dc.border = border_all
        ws.row_dimensions[current_row].height = 85

        for i, p in enumerate(periods):
            col = i + 2
            cell_entries = timetable_dict.get((day, p), [])
            if not cell_entries:
                cell_text = "Empty Slot"
                fill = fill_empty
                cell_font = Font(italic=True, color="9CA3AF", size=10)
            else:
                is_lab = any(e.session_type and e.session_type.upper() == 'LAB' for e in cell_entries)
                lines = []
                for e in cell_entries:
                    block = f"{e.course_code}\n{e.course_name or ''}\n({e.faculty_name or 'Unassigned'})\n[{e.venue_name or ''}]"
                    lines.append(block)
                cell_text = "\n---\n".join(lines)
                fill = fill_lab if is_lab else fill_class
                cell_font = Font(bold=False, size=10, color="111827")

            c = ws.cell(row=current_row, column=col, value=cell_text)
            c.alignment = align_center
            c.fill = fill
            c.border = border_all
            c.font = cell_font
        current_row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
